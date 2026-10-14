from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "FieldNode_Proposta_TCC_v5.docx"


def set_font(run, name="Calibri", size=11, color="24313A", bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for item in list(grid):
        grid.remove(item)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl_pr = table._tbl.tblPr
    width_node = tbl_pr.first_child_found_in("w:tblW")
    if width_node is None:
        width_node = OxmlElement("w:tblW")
        tbl_pr.append(width_node)
    width_node.set(qn("w:w"), str(sum(widths)))
    width_node.set(qn("w:type"), "dxa")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_para(doc, text="", style=None, align=None, after=8, before=0, line=1.333):
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if text:
        set_font(paragraph.add_run(text))
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    set_font(paragraph.add_run(text))
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    set_font(paragraph.add_run(text), size={1: 16, 2: 13, 3: 12}[level], color="2E74B5" if level < 3 else "1F4D78", bold=True)
    return paragraph


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("24313A")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for level, size, color, before, after in ((1, 16, "2E74B5", 18, 10), (2, 13, "2E74B5", 12, 6), (3, 12, "1F4D78", 8, 4)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("FieldNode | Proposta de TCC v5"), size=9, color="68747B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Documento revisado em 12/10/2026"), size=9, color="68747B")

    add_para(doc, "FIELDNODE", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    title = add_para(doc, "Telemetria offline-first e inteligência operacional para máquinas agrícolas", align=WD_ALIGN_PARAGRAPH.CENTER, after=6, line=1.15)
    for run in title.runs:
        set_font(run, size=24, color="173B3F", bold=True)
    subtitle = add_para(doc, "Proposta revisada com base no estado implementado até a Semana 10", align=WD_ALIGN_PARAGRAPH.CENTER, after=22, line=1.15)
    for run in subtitle.runs:
        set_font(run, size=13, color="68747B", italic=True)

    metadata = doc.add_table(rows=4, cols=2)
    set_table_widths(metadata, [2700, 6660])
    rows = [("Natureza", "Protótipo técnico funcional para TCC"), ("Escopo demonstrável", "Leitura -> Trust Score -> Event -> Decision -> aprovação humana -> outcome"), ("Base técnica", "Django/DRF, SQLite/MySQL, Next.js, cliente de borda com outbox"), ("Limite declarado", "Não demonstra sensor físico calibrado, economia medida ou manutenção executada automaticamente")]
    for row, (label, value) in zip(metadata.rows, rows):
        shade(row.cells[0], "E8EEF5")
        for cell, text in zip(row.cells, (label, value)):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(text), size=10, bold=cell is row.cells[0])

    add_heading(doc, "1. Resumo", 1)
    add_para(doc, "O FieldNode é uma plataforma de telemetria e inteligência operacional para máquinas agrícolas. A proposta combina ingestão autenticada, validação, deduplicação, persistência, indicadores de qualidade dos dados, detecção de eventos e recomendações operacionais submetidas à aprovação humana.")
    add_para(doc, "A revisão desta proposta mantém a ambição do produto, mas ajusta as afirmações ao que foi realmente implementado. O resultado atual é um protótipo técnico funcional, demonstrável em ambiente controlado e ainda dependente de validação em campo e de repetição oficial no Docker Compose.")

    add_heading(doc, "2. Problema e justificativa", 1)
    add_para(doc, "Máquinas agrícolas operam em ambientes com conectividade irregular, múltiplas fontes de dados e necessidade de decisão rápida. Um sistema que apenas exibe leituras não resolve o problema operacional: é necessário preservar a origem, reconhecer condições relevantes, explicar a recomendação e registrar quem decidiu o quê.")
    add_para(doc, "A justificativa técnica está na combinação de integridade da telemetria, tolerância a falhas de comunicação e memória auditável de decisões. A justificativa científica está em medir o comportamento do protótipo com critérios explícitos, sem apresentar regras ou sinais de aprendizado não supervisionado como predição supervisionada validada.")

    add_heading(doc, "3. Objetivos", 1)
    add_heading(doc, "3.1 Objetivo geral", 2)
    add_para(doc, "Projetar e demonstrar um protótipo de telemetria offline-first capaz de transformar leituras de máquinas agrícolas em eventos e recomendações operacionais auditáveis, com revisão humana antes de qualquer registro de execução.")
    add_heading(doc, "3.2 Objetivos específicos", 2)
    for item in [
        "receber telemetria por HTTP e preservar identidade de dispositivo, mensagem, sequência, timestamp e hash;",
        "aplicar validação, deduplicação e uma outbox local com retry e cursor contíguo;",
        "calcular e persistir Trust Score heurístico e estado agregado de qualidade;",
        "gerar Event determinístico para condições críticas documentadas e associá-lo à leitura de origem;",
        "gerar Decision persistente, vinculada ao Event quando houver evento aberto elegível;",
        "permitir aprovação ou rejeição por usuário autenticado e registrar outcome textual;",
        "documentar limites, evidências e pendências para evolução posterior."
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. Arquitetura e fluxo", 1)
    add_para(doc, "A arquitetura atual é composta por cliente de borda ou simulador, API Django/DRF, banco SQLite ou MySQL, pipeline de análise e dashboard Next.js. A ingestão usa X-API-Key para o canal máquina-a-máquina. A aprovação usa autenticação de usuário e autorização no backend.")
    add_para(doc, "No fluxo fechado da Semana 10, uma leitura determinística de 115 °C é enviada pelo EdgeClient. A API calcula o Trust Score, persiste a leitura, avalia a regra de temperatura acima de 85 °C, cria um Event TEMP_ALTA crítico e persiste uma Decision pendente associada ao Event. O operador revisa a recomendação e aprova pela interface. O outcome deve descrever o que ocorreu, sem afirmar execução física inexistente.")

    add_heading(doc, "5. Metodologia", 1)
    add_para(doc, "A avaliação combina inspeção de contratos, testes automatizados, execução local e ensaio integrado no Compose. O critério de qualificação do closed loop exige três execuções Compose consecutivas, com persistência, rastreabilidade, aprovação humana e outcome confirmados. Uma execução local SQLite é evidência parcial e não substitui o ambiente oficial.")
    add_para(doc, "A análise de risco combina regras determinísticas e Isolation Forest como sinais complementares. A PoC de decisão agentiva é mantida como experimento condicionado a chave externa, Event elegível e resposta auditável; não é requisito para afirmar que o closed loop básico funciona.")

    add_heading(doc, "6. Resultados verificáveis até a Semana 10", 1)
    for item in [
        "ingestão autenticada, validação e deduplicação por chave composta no backend;",
        "outbox SQLite do EdgeClient com ACK, hash, retry e cursor;",
        "Trust Score individual e MachineDataHealth persistidos;",
        "Event crítico por temperatura e Decision pendente vinculada ao Event;",
        "transições de Decision protegidas por estado, usuário autenticado e autorização da Machine;",
        "dashboard com revisão e registro de outcome;",
        "testes focados de Trust Score, Event, Decision e telemetria executados sem falhas;",
        "roteiro de banca e respostas de pressão técnica documentados."
    ]:
        add_bullet(doc, item)

    add_heading(doc, "7. Limitações e trabalho futuro", 1)
    for item in [
        "a leitura demonstrada é simulada; a integração física com sensor e a calibração ainda precisam de validação;",
        "Trust Score e regras atuais são heurísticos e não equivalem a probabilidade calibrada de falha;",
        "a execução física da intervenção não é controlada pelo protótipo;",
        "a integração do modelo universal de telemetria com todo o pipeline ainda é uma dívida técnica;",
        "a autenticação do MVP requer evolução antes de uso produtivo, incluindo sessão adequada, rotação e controles adicionais;",
        "a validação comercial com compradores e usuários operacionais permanece exploratória;",
        "a execução oficial dos três ensaios Compose e a gravação do vídeo dependem de Docker Desktop disponível."
    ]:
        add_bullet(doc, item)

    add_heading(doc, "8. Contribuição esperada", 1)
    add_para(doc, "A contribuição do trabalho é uma demonstração integrada de como telemetria offline-first, qualidade dos dados, regras de eventos, sinais de análise e revisão humana podem formar um ciclo auditável para operação agrícola. O valor da proposta está na rastreabilidade e na clareza dos limites, não em prometer autonomia ou resultados de campo que ainda não foram medidos.")

    add_heading(doc, "9. Referências internas do projeto", 1)
    for item in [
        "docs/ROTEIRO_DEMO_BANCA.md",
        "docs/INVESTOR_CTO_PRESSURE_TEST.md",
        "docs/S10_EXECUCOES.md",
        "docs/HONESTIDADE_METODOLOGICA_ML.md",
        "docs/CONTRATO_TELEMETRIA_V1.md",
        "docs/SEGURANCA.md",
    ]:
        add_bullet(doc, item)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
